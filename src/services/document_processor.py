"""Document processing service for loading and parsing various file formats.

Implements a multi-step pipeline per document:
1) Extract → 2) Clean → 3) Validate → 4) Analyze (external) → 5) Summarize (external).
"""
import os
from datetime import datetime
from pathlib import Path
from typing import List

from pypdf import PdfReader
from docx import Document as DocxDocument

from src.models.document import Document
from src.utils.logger import get_logger
from config.settings import SUPPORTED_FORMATS

logger = get_logger(__name__)


class UnsupportedFormatError(Exception):
    """Raised when file format is not supported."""

    pass


class ParseError(Exception):
    """Raised when file parsing fails."""

    pass


class DocumentProcessor:
    """Handles loading and parsing documents from various file formats."""

    def load_document(self, file_path: str, fallback_index: int | None = None) -> Document:
        """
        Load and parse a single document file.

        Args:
            file_path: Absolute path to document file

        Returns:
            Document object with extracted content

        Raises:
            FileNotFoundError: If file doesn't exist
            UnsupportedFormatError: If file format not supported
            ParseError: If file parsing fails
        """
        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        file_format = path.suffix.lower()
        if file_format not in SUPPORTED_FORMATS:
            raise UnsupportedFormatError(
                f"Unsupported format: {file_format}. "
                f"Supported formats: {SUPPORTED_FORMATS}"
            )

        try:
            # Step 1/2: extract and clean text based on file format
            if file_format in [".txt", ".md"]:
                raw_content = self._parse_text_file(path)
            elif file_format == ".pdf":
                raw_content = self._parse_pdf(path)
            elif file_format == ".docx":
                raw_content = self._parse_docx(path)
            else:
                raise UnsupportedFormatError(f"Unsupported format: {file_format}")

            content = self._clean_text(raw_content)

            # Step 3: basic validation on extracted text length
            if len(content) < 300:
                logger.warning(
                    "Extracted content too short for %s (len=%d), attempting basic re-parse",
                    path.name,
                    len(content),
                )
                # For now, we simply retry once for PDFs and DOCX; more
                # advanced strategies (page batching, alternate libraries)
                # can be added here.
                if file_format == ".pdf":
                    raw_content = self._parse_pdf(path)
                elif file_format == ".docx":
                    raw_content = self._parse_docx(path)
                content = self._clean_text(raw_content)

            # Derive a robust filename with optional fallback
            filename = path.name if path.name else ""
            if not filename and fallback_index is not None:
                filename = f"Document-{fallback_index}"
            elif not filename:
                filename = "Document-unknown"

            # Create Document object
            stat = path.stat()
            document = Document(
                file_path=str(path.absolute()),
                filename=filename,
                format=file_format,
                size_bytes=stat.st_size,
                content=content,
                created_at=datetime.fromtimestamp(stat.st_ctime),
                status="pending",
            )

            # Detect and attach a best-effort title for later use in prompts
            document.title = self.detect_title(document)

            logger.info(
                "Loaded document: %s (size=%d bytes, extracted=%d chars)",
                path.name,
                stat.st_size,
                len(content),
            )
            return document

        except Exception as e:
            logger.error(f"Failed to parse {file_path}: {e}", exc_info=True)
            raise ParseError(f"Failed to parse {file_path}: {e}")

    def _parse_text_file(self, path: Path) -> str:
        """Parse plain text file (.txt, .md)."""
        try:
            return path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            # Fallback to latin-1 encoding
            logger.warning(f"UTF-8 decode failed for {path}, trying latin-1")
            return path.read_text(encoding="latin-1")

    def _parse_pdf(self, path: Path) -> str:
        """Parse PDF file."""
        try:
            reader = PdfReader(str(path))
            text_parts = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)
            content = "\n".join(text_parts)
            return content.strip()
        except Exception as e:
            raise ParseError(f"PDF parsing error: {e}")

    def _parse_docx(self, path: Path) -> str:
        """Parse DOCX file."""
        try:
            doc = DocxDocument(str(path))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            content = "\n".join(paragraphs)
            return content.strip()
        except Exception as e:
            raise ParseError(f"DOCX parsing error: {e}")

    def _clean_text(self, text: str) -> str:
        """Normalize extracted text.

        - Remove broken Unicode replacement chars
        - Merge hyphenated line breaks
        - Collapse excessive empty lines
        """
        if not text:
            return ""

        cleaned = text.replace("\uFFFD", "")

        # Merge hyphenated line breaks: "some-
        # thing" -> "something"
        cleaned = cleaned.replace("-\n", "")

        # Normalize line endings and collapse multiple blank lines
        lines = cleaned.splitlines()
        normalized_lines: List[str] = []
        empty_run = 0
        for line in lines:
            stripped = line.rstrip()
            if not stripped:
                empty_run += 1
                if empty_run > 1:
                    continue
            else:
                empty_run = 0
            normalized_lines.append(stripped)

        return "\n".join(normalized_lines).strip()

    def detect_title(self, document: Document) -> str:
        """Best-effort title detection using simple heuristics.

        - Prefer first non-empty line from the content
        - Fallback to filename without extension
        - If all else fails, synthesize a generic title
        """
        # Try first non-empty line of content
        for line in document.content.splitlines():
            candidate = line.strip()
            if candidate:
                return candidate[:120]

        # Fallback to filename stem
        path = Path(document.file_path)
        if path.stem:
            return path.stem

        return "Untitled Document"

    def load_documents_from_folder(self, folder_path: str) -> List[Document]:
        """
        Load all supported documents from a folder (recursive scan).

        Args:
            folder_path: Path to folder containing documents

        Returns:
            List of Document objects (both successful and failed)
        """
        folder = Path(folder_path)
        if not folder.exists():
            logger.error(f"Folder not found: {folder_path}")
            return []

        documents: List[Document] = []
        file_count = 0
        fallback_counter = 1

        # Recursively scan for supported files
        for file_path in folder.rglob("*"):
            if file_path.is_file() and file_path.suffix.lower() in SUPPORTED_FORMATS:
                file_count += 1
                try:
                    document = self.load_document(str(file_path), fallback_index=fallback_counter)
                    documents.append(document)
                    fallback_counter += 1
                except Exception as e:
                    # Create failed document entry
                    logger.error(f"Failed to load {file_path.name}: {e}")
                    failed_doc = Document(
                        file_path=str(file_path.absolute()),
                        filename=file_path.name,
                        format=file_path.suffix.lower(),
                        size_bytes=file_path.stat().st_size if file_path.exists() else 0,
                        content="",
                        status="failed",
                        error_message=str(e),
                    )
                    documents.append(failed_doc)

        logger.info(
            f"Scanned {folder_path}: found {file_count} files, "
            f"loaded {sum(1 for d in documents if d.status != 'failed')} successfully"
        )
        return documents
